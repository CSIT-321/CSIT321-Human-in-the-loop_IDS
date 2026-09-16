#!/usr/bin/env python3
"""Collision check for the deck.

python-pptx text boxes auto-grow downwards in PowerPoint, so a box declared
too short does not clip -- it overlaps whatever is below it. This script
estimates, for every text frame, the height it really needs and the width its
text really occupies (respecting alignment), then reports only the frames
whose real extent reaches another shape or runs off the slide.
"""
from __future__ import annotations

import math
import sys

from pptx import Presentation
from pptx.enum.text import PP_ALIGN

EMU_IN = 914400.0
SAFE_BOTTOM = 7.28      # inches


def para_stats(p):
    text = "".join(r.text for r in p.runs)
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    size = max(sizes) if sizes else 18.0
    mono = any(r.font.name and "consol" in r.font.name.lower() for r in p.runs)
    factor = 0.55 if mono else 0.485
    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
    sb = p.space_before.pt if p.space_before else 0.0
    sa = p.space_after.pt if p.space_after else 0.0
    return text, size, factor, ls, sb, sa


def frame_metrics(tf, usable_w_pt):
    """(needed_height_pt, widest_line_pt, align)"""
    needed = 0.0
    widest = 0.0
    align = PP_ALIGN.LEFT
    for p in tf.paragraphs:
        text, size, factor, ls, sb, sa = para_stats(p)
        if p.alignment is not None:
            align = p.alignment
        cpl = max(usable_w_pt / (factor * size), 4.0)
        chunks = text.split("\n") if text else [""]
        lines = 0
        for chunk in chunks:
            n = max(1, math.ceil(len(chunk) / cpl))
            lines += n
            # the last line of a chunk is short unless it wrapped exactly
            full = math.floor(len(chunk) / cpl) if len(chunk) else 0
            if full:
                widest = max(widest, usable_w_pt)
            tail = len(chunk) - full * cpl
            if tail > 0:
                widest = max(widest, tail * factor * size)
        needed += lines * max(size * 1.22 * (ls or 1.0), 11.0) + sb + sa
    return needed, min(widest, usable_w_pt), align


def boxes(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_table:
            out.append(("table", shape.left / EMU_IN, shape.top / EMU_IN,
                        shape.width / EMU_IN, shape.height / EMU_IN, None))
            continue
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        tf = shape.text_frame
        ml = (tf.margin_left or 0) / EMU_IN * 72
        mr = (tf.margin_right or 0) / EMU_IN * 72
        usable = shape.width / EMU_IN * 72 - ml - mr
        need_pt, wide_pt, align = frame_metrics(tf, max(usable, 12))
        x0 = shape.left / EMU_IN
        w = shape.width / EMU_IN
        used_w = min(max(wide_pt / 72.0, 0.10), w)
        if align == PP_ALIGN.RIGHT:
            x0 = x0 + w - used_w
        elif align == PP_ALIGN.CENTER:
            x0 = x0 + (w - used_w) / 2.0
        out.append(("text", x0, shape.top / EMU_IN, used_w,
                    need_pt / 72.0, tf.text))
    return out


def check(path: str) -> int:
    prs = Presentation(path)
    found = 0
    for idx, slide in enumerate(prs.slides, start=1):
        items = boxes(slide)
        for i, (kind, left, top, width, height, text) in enumerate(items):
            if kind != "text":
                continue
            bottom = top + height
            issues = []
            if bottom > SAFE_BOTTOM:
                issues.append(f"runs past safe bottom ({bottom:.2f}in)")
            for j, (k2, l2, t2, w2, h2, txt2) in enumerate(items):
                if i == j:
                    continue
                if abs(t2 - top) < 0.02 and abs(l2 - left) < 0.02:
                    continue
                if not ((left + 0.03 < l2 + w2) and (l2 + 0.03 < left + width)):
                    continue
                if t2 > top + 0.06 and bottom > t2 + 0.03:
                    snippet = (txt2 or "").strip().replace("\n", " ")[:32]
                    issues.append(f"reaches {k2} at y={t2:.2f} ({snippet})")
            if issues:
                found += 1
                print(f"slide {idx:>2}  {text.strip()[:44]!r} ends "
                      f"{bottom:5.2f}in")
                for issue in issues:
                    print(f"           -> {issue}")
    print("\nno collisions found" if not found else f"\n{found} problem frames")
    return 1 if found else 0


if __name__ == "__main__":
    sys.exit(check(sys.argv[1]))
