"""
Build the Preliminary User Manual as a Word document.

    python scripts/build_user_manual.py

**One source of truth.** The prose lives in `docs/preliminary-user-manual.md`; this script renders
it. Editing the Word file directly would fork the two, and this project has already spent a session
repairing documents that drifted apart. Change the Markdown, rebuild, hand in the `.docx`.

The layout follows the sample manual the school supplied:

    page 1      cover - school, unit, title, project, group, assessor, supervisor, team table
    page 2      document control - title, document name, record of revision
    page 3      table of contents (a real Word field, so the page numbers are correct)
    page 4+     sections 1-4, with a running header and a "Page N | M" footer

**Details we do not have are written as visible placeholders**, never invented and never left
silently blank: an assessor's name guessed wrong is worse than an obvious [TO BE COMPLETED].

Output: hitl-ids/docs/FYP-26-S3-13_PrelimUserManual.docx, or the path given with --out:

    python scripts/build_user_manual.py --out docs/FYP-26-S3-13_PrelimUserManual_v0.2.docx

Use --out whenever the default file holds edits made in Word: a rebuild replaces it wholesale.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HITL = Path(__file__).resolve().parent.parent
SOURCE = HITL / "docs" / "preliminary-user-manual.md"
TARGET = HITL / "docs" / "FYP-26-S3-13_PrelimUserManual.docx"

#: Cover-page details. Square brackets mark what the team still has to fill in.
COVER = {
    "school": "School of Computing and Information Technology",
    "unit": "CSIT321 - Project",
    "title": "User Manual",
    "topic_code": "[CSIT-26-S3-XX]",
    "topic_name": "(Human-in-the-Loop Intrusion Detection Dashboard)",
    "group": "FYP-26-S3-13",
    "assessor": "[TO BE COMPLETED]",
    "supervisor": "[TO BE COMPLETED]",
}

#: One row per team member. Only the author is known here; the rest are for the team to complete.
TEAM = [
    ("Glenn Ang", "[UOW ID]", "[SIM email]"),
    ("[Name]", "[UOW ID]", "[SIM email]"),
    ("[Name]", "[UOW ID]", "[SIM email]"),
    ("[Name]", "[UOW ID]", "[SIM email]"),
    ("[Name]", "[UOW ID]", "[SIM email]"),
    ("[Name]", "[UOW ID]", "[SIM email]"),
]

DOC_TITLE = "User Manual"
DOC_NAME = "FYP-26-S3-13 User Manual, Version 0.2"

INK = RGBColor(0x1F, 0x1F, 0x1F)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
QUIET = RGBColor(0x59, 0x59, 0x59)
CODE_BG = "F2F4F7"
CALLOUT_BG = "EAF1F8"


# --------------------------------------------------------------------------------------------
# Word plumbing that python-docx does not wrap
# --------------------------------------------------------------------------------------------

def text_element(value: str):
    node = OxmlElement("w:t")
    node.text = value
    return node


def field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, NUMPAGES, TOC). Word computes the value when it opens."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate):
        run._r.append(element)
    if placeholder:
        run._r.append(text_element(placeholder))
    run._r.append(end)


def update_fields_on_open(document: Document) -> None:
    """Ask Word to refresh every field when the document opens.

    Without this the table of contents shows its placeholder until somebody presses F9, which reads
    as an unfinished document rather than a generated one.
    """
    flag = OxmlElement("w:updateFields")
    flag.set(qn("w:val"), "true")
    document.settings.element.append(flag)


def shade(element, colour: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), colour)
    element.append(shading)


def cell_shade(cell, colour: str) -> None:
    shade(cell._tc.get_or_add_tcPr(), colour)


# --------------------------------------------------------------------------------------------
# A very small Markdown reader - only the constructs this manual actually uses
# --------------------------------------------------------------------------------------------

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
IMAGE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
HEADING = re.compile(r"^(?P<hashes>#{1,4})\s+(?P<text>.+?)\s*$")
BULLET = re.compile(r"^[-*]\s+(?P<text>.+?)\s*$")
NUMBERED = re.compile(r"^(?P<n>\d+)\.\s+(?P<text>.+?)\s*$")

#: Headings the front matter lays out by hand, so the body must not repeat them.
FRONT_MATTER = {"preliminary user manual", "document control", "record of revision"}


def write_runs(paragraph, source: str, *, size: float = 10.5, colour=INK,
               italic: bool = False) -> None:
    """Render inline **bold**, *italic* and `code`, and reduce a link to its text."""
    source = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", source)
    for piece in INLINE.split(source):
        if not piece:
            continue
        run = paragraph.add_run()
        if piece.startswith("**") and piece.endswith("**"):
            run.text, run.bold = piece[2:-2], True
        elif piece.startswith("`") and piece.endswith("`"):
            run.text = piece[1:-1]
            run.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*"):
            run.text, run.italic = piece[1:-1], True
        else:
            run.text = piece
        run.font.size = Pt(size)
        run.font.color.rgb = colour
        if italic:
            run.italic = True


def add_table(document: Document, rows: list[list[str]], width: float) -> None:
    if not rows:
        return
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, header):
        cell.text = ""
        write_runs(cell.paragraphs[0], f"**{label}**", size=9.5)
        cell_shade(cell, "DDE6F0")
    for line in body:
        cells = table.add_row().cells
        for cell, value in zip(cells, line):
            cell.text = ""
            write_runs(cell.paragraphs[0], value, size=9.5)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(width / len(header))
    document.add_paragraph()


def add_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    shade(paragraph._p.get_or_add_pPr(), CODE_BG)
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if i < len(lines) - 1:
            run.add_break()


def add_callout(document: Document, lines: list[str], width: float) -> None:
    """A quoted aside, as a single shaded cell so it reads as a box rather than an indent."""
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.width = Inches(width)
    cell_shade(cell, CALLOUT_BG)
    cell.text = ""
    for i, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        write_runs(paragraph, line, size=10)
    document.add_paragraph()


def add_figure(document: Document, src: Path, alt: str, width: float) -> None:
    if not src.exists():
        raise SystemExit(f"missing figure {src}\n"
                         "screenshots: GUIDE_CAPTURE=1 npx playwright test capture-guide (in apps/web)\n"
                         "wireframes:  python scripts/make_wireframes.py")
    document.add_picture(str(src), width=Inches(width))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    write_runs(caption, alt, size=9, colour=QUIET, italic=True)


def render_body(document: Document, markdown: str, width: float) -> None:
    """Walk the Markdown and emit Word. Front-matter sections are skipped."""
    lines = markdown.splitlines()
    index, skipping = 0, False
    pending: list[list[str]] = []

    def flush() -> None:
        nonlocal pending
        if pending:
            add_table(document, pending, width)
            pending = []

    while index < len(lines):
        stripped = lines[index].strip()

        heading = HEADING.match(stripped)
        if heading:
            flush()
            title, level = heading.group("text"), len(heading.group("hashes"))
            skipping = level <= 3 and title.lower() in FRONT_MATTER
            if not skipping:
                if level == 2:
                    document.add_page_break()
                paragraph = document.add_heading(level=min(level - 1, 3))
                write_runs(paragraph, title, size=16 - 2 * (level - 2), colour=ACCENT)
            index += 1
            continue
        if skipping:
            index += 1
            continue

        if not stripped or stripped == "---":
            flush()
            index += 1
            continue

        picture = IMAGE.match(stripped)
        if picture:
            flush()
            add_figure(document, (SOURCE.parent / picture.group("src")).resolve(),
                       picture.group("alt"), width)
            index += 1
            continue

        if stripped.startswith("```"):
            flush()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(document, block)
            index += 1
            continue

        if stripped.startswith(">"):
            flush()
            block = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_callout(document, [b for b in block if b], width)
            continue

        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):  # the |---|---| separator row
                index += 1
                continue
            pending.append(cells)
            index += 1
            continue
        flush()

        bullet = BULLET.match(stripped)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            write_runs(paragraph, bullet.group("text"))
            index += 1
            continue

        numbered = NUMBERED.match(stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            write_runs(paragraph, numbered.group("text"))
            index += 1
            continue

        # An ordinary paragraph: rejoin the lines the Markdown wrapped.
        block = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "- ", "* "))
                    or IMAGE.match(nxt) or NUMBERED.match(nxt) or nxt == "---"):
                break
            block.append(nxt)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        write_runs(paragraph, " ".join(block))
    flush()


# --------------------------------------------------------------------------------------------
# The front matter, laid out to match the sample
# --------------------------------------------------------------------------------------------

def build_cover(document: Document, width: float) -> None:
    document.add_paragraph()
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    write_runs(logo, "[SCHOOL LOGO - paste from the supplied template]", size=9, colour=QUIET,
               italic=True)

    for value, size, bold in ((COVER["school"], 16, True), (COVER["unit"], 14, False),
                              (COVER["title"], 26, True)):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(10)
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = ACCENT if bold else INK

    document.add_paragraph()
    for label, value in (("Project Topic", COVER["topic_code"]), ("", COVER["topic_name"]),
                         ("Group Number", COVER["group"]), ("Assessor", COVER["assessor"]),
                         ("Supervisor", COVER["supervisor"])):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        write_runs(paragraph, f"**{label}:** {value}" if label else value, size=12)

    document.add_paragraph()
    add_table(document, [["Name", "UOW ID", "SIM Email"], *[list(r) for r in TEAM]], width)


def revision_rows(markdown: str) -> list[list[str]]:
    """Lift the revision table out of the Markdown, so it is maintained in one place."""
    parts = markdown.split("### Record of revision", 1)
    if len(parts) < 2:
        return [["Date", "Description", "Section affected", "Changes made by",
                 "Version after revision"]]
    rows: list[list[str]] = []
    for line in parts[1].splitlines():
        line = line.strip()
        if not line.startswith("|"):
            if rows:
                break
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if set("".join(cells)) <= set("-: "):
            continue
        rows.append(cells)
    return rows


def build_document_control(document: Document, markdown: str, width: float) -> None:
    document.add_page_break()
    heading = document.add_heading(level=1)
    write_runs(heading, "Document Control", size=16, colour=ACCENT)
    for label, value in (("Title", DOC_TITLE), ("Document Name", DOC_NAME)):
        paragraph = document.add_paragraph()
        write_runs(paragraph, f"**{label}:** {value}", size=11)

    heading = document.add_heading(level=2)
    write_runs(heading, "Record of Revision", size=13, colour=ACCENT)
    add_table(document, revision_rows(markdown), width)


def build_contents(document: Document) -> None:
    document.add_page_break()
    heading = document.add_heading(level=1)
    write_runs(heading, "Table of Contents", size=16, colour=ACCENT)
    paragraph = document.add_paragraph()
    field(paragraph, r'TOC \o "1-3" \h \z \u',
          "Right-click and choose Update Field if this does not fill in automatically.")


def build_chrome(document: Document) -> None:
    """The running header and the 'Page N | M' footer the sample uses."""
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    write_runs(header, DOC_TITLE, size=9, colour=QUIET)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    field(footer, "PAGE", "1")
    footer.add_run(" | ")
    field(footer, "NUMPAGES", "1")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = QUIET


def main() -> int:
    parser = argparse.ArgumentParser(description="Render docs/preliminary-user-manual.md as .docx")
    parser.add_argument("--out", type=Path, default=TARGET,
                        help="where to write the .docx (default: %(default)s)")
    target = parser.parse_args().out
    if not target.is_absolute():
        target = HITL / target
    if not SOURCE.exists():
        print(f"missing {SOURCE}", file=sys.stderr)
        return 2
    markdown = SOURCE.read_text(encoding="utf-8")

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)  # A4
    section.left_margin = section.right_margin = Inches(1.0)
    section.top_margin = section.bottom_margin = Inches(0.9)
    width = 8.27 - 2.0

    build_chrome(document)
    build_cover(document, width)
    build_document_control(document, markdown, width)
    build_contents(document)
    render_body(document, markdown, width)
    update_fields_on_open(document)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    # Python 3.11 forbids a backslash inside an f-string expression, so count first.
    figures = len(re.findall(r"^!\[", markdown, re.M))
    print(f"wrote {target}")
    print(f"  {figures} figures, {len(document.paragraphs)} paragraphs, "
          f"{len(document.tables)} tables")
    print("  open it in Word once so the contents and page numbers fill in")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
